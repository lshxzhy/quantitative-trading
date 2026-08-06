"""Maintain the flat Tushare Pickle bundle and its Excel viewing copy."""

from __future__ import annotations

import argparse
import os
from datetime import datetime, time, timedelta, timezone
from pathlib import Path
from typing import Any
from uuid import uuid4
from zoneinfo import ZoneInfo

import pandas as pd
import tushare as ts
from dotenv import load_dotenv
from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .tushare_pipeline import (
    ENDPOINT_SPECS,
    INDEX_COLUMNS,
    STOCK_COLUMNS,
    TushareApi,
    build_indexes_frame,
    build_stocks_frame,
    fetch_trade_date,
    parse_trade_date,
    query_frame,
    validate_tables,
)


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_DATA_ROOT = PROJECT_ROOT / "data"
CHINA_TIMEZONE = ZoneInfo("Asia/Shanghai")
DATA_READY_TIME = time(18, 0)
BUNDLE_SCHEMA_VERSION = 1
BUNDLE_SOURCE = "tushare_pro"
EXCEL_MAX_DATA_ROWS = 1_048_575


def default_end_date(now: datetime | None = None) -> str:
    current = now or datetime.now(CHINA_TIMEZONE)
    if current.tzinfo is None:
        raise ValueError("now must include timezone information")

    china_now = current.astimezone(CHINA_TIMEZONE)
    candidate = china_now.date()
    if china_now.time() < DATA_READY_TIME:
        candidate -= timedelta(days=1)
    return candidate.strftime("%Y%m%d")


def missing_trade_dates(open_dates: list[str], existing_dates: list[str]) -> list[str]:
    existing = set(existing_dates)
    return [trade_date for trade_date in open_dates if trade_date not in existing]


def query_open_dates(api: TushareApi, start_date: str, end_date: str) -> list[str]:
    spec = ENDPOINT_SPECS["trade_cal"]
    frame = query_frame(
        api,
        "trade_cal",
        spec,
        exchange="SSE",
        start_date=start_date,
        end_date=end_date,
    )
    if frame.empty:
        raise RuntimeError(
            f"Tushare trade_cal returned no rows for {start_date} through {end_date}"
        )
    if frame.duplicated(list(spec.keys)).any():
        raise RuntimeError("Tushare trade_cal returned duplicate exchange/date rows")

    actual_dates = frame["cal_date"].astype(str)
    outside = actual_dates[(actual_dates < start_date) | (actual_dates > end_date)]
    if not outside.empty:
        raise RuntimeError(
            "Tushare trade_cal returned dates outside the requested range: "
            + ", ".join(sorted(set(outside)))
        )

    open_dates = frame.loc[frame["is_open"].astype(str).eq("1"), "cal_date"]
    return sorted(open_dates.astype(str).tolist())


def bundle_paths(data_root: Path) -> tuple[Path, Path]:
    root = data_root.resolve()
    return root / "market_data.pkl", root / "market_data.xlsx"


def make_bundle(stocks: pd.DataFrame, indexes: pd.DataFrame) -> dict[str, Any]:
    validate_tables(stocks, indexes)
    return {
        "schema_version": BUNDLE_SCHEMA_VERSION,
        "source": BUNDLE_SOURCE,
        "updated_at_utc": datetime.now(timezone.utc).isoformat(),
        "stocks": stocks,
        "indexes": indexes,
    }


def validate_bundle(bundle: object) -> dict[str, Any]:
    if not isinstance(bundle, dict):
        raise RuntimeError("market_data.pkl must contain a dictionary")
    expected_keys = {"schema_version", "source", "updated_at_utc", "stocks", "indexes"}
    if set(bundle) != expected_keys:
        raise RuntimeError(
            f"market_data.pkl keys are {sorted(bundle)}, expected {sorted(expected_keys)}"
        )
    if bundle["schema_version"] != BUNDLE_SCHEMA_VERSION:
        raise RuntimeError(
            f"Unsupported market_data.pkl schema version: {bundle['schema_version']}"
        )
    if bundle["source"] != BUNDLE_SOURCE:
        raise RuntimeError(f"Unexpected market_data.pkl source: {bundle['source']}")
    try:
        parsed_updated_at = datetime.fromisoformat(str(bundle["updated_at_utc"]))
    except ValueError as error:
        raise RuntimeError("updated_at_utc is not a valid ISO timestamp") from error
    if parsed_updated_at.tzinfo is None:
        raise RuntimeError("updated_at_utc must include a UTC offset")
    if not isinstance(bundle["stocks"], pd.DataFrame):
        raise RuntimeError("market_data.pkl stocks value is not a DataFrame")
    if not isinstance(bundle["indexes"], pd.DataFrame):
        raise RuntimeError("market_data.pkl indexes value is not a DataFrame")
    validate_tables(bundle["stocks"], bundle["indexes"])
    return bundle


def load_bundle(path: Path) -> dict[str, Any]:
    if not path.is_file():
        raise FileNotFoundError(f"Tushare bundle does not exist: {path}")
    return validate_bundle(pd.read_pickle(path))


def _complete_recent_dates(frame: pd.DataFrame, maximum_rows: int) -> list[str]:
    counts = frame.groupby("trade_date", sort=True).size()
    selected: list[str] = []
    row_count = 0
    for trade_date, count in reversed(list(counts.items())):
        count = int(count)
        if row_count + count > maximum_rows:
            break
        selected.append(str(trade_date))
        row_count += count
    if not selected:
        raise RuntimeError("The newest complete trading day exceeds Excel's row limit")
    return sorted(selected)


def excel_view_frames(
    stocks: pd.DataFrame, indexes: pd.DataFrame
) -> tuple[pd.DataFrame, pd.DataFrame]:
    stock_dates = _complete_recent_dates(stocks, EXCEL_MAX_DATA_ROWS)
    index_dates = _complete_recent_dates(indexes, EXCEL_MAX_DATA_ROWS)
    stock_view = stocks.loc[stocks["trade_date"].astype(str).isin(stock_dates)]
    index_view = indexes.loc[indexes["trade_date"].astype(str).isin(index_dates)]
    return stock_view, index_view


def _excel_value(value: object) -> object:
    if pd.isna(value):
        return None
    item = getattr(value, "item", None)
    return item() if callable(item) else value


def _set_column_widths(worksheet: object, columns: tuple[str, ...]) -> None:
    wide = {"fullname": 28, "enname": 36, "industry": 18, "name": 16}
    medium = {
        "trade_date": 12,
        "ts_code": 13,
        "list_date": 12,
        "delist_date": 12,
        "suspend_timing": 20,
        "suspend_type": 20,
    }
    for position, column in enumerate(columns, start=1):
        width = wide.get(column, medium.get(column, max(11, min(len(column) + 2, 18))))
        worksheet.column_dimensions[get_column_letter(position)].width = width


def _append_sheet(
    workbook: Workbook,
    name: str,
    frame: pd.DataFrame,
    columns: tuple[str, ...],
) -> None:
    worksheet = workbook.create_sheet(name)
    worksheet.freeze_panes = "A2"
    _set_column_widths(worksheet, columns)

    fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    font = Font(color="FFFFFF", bold=True)
    header = []
    for value in columns:
        cell = WriteOnlyCell(worksheet, value=value)
        cell.fill = fill
        cell.font = font
        header.append(cell)
    worksheet.append(header)

    for row in frame.loc[:, columns].itertuples(index=False, name=None):
        worksheet.append([_excel_value(value) for value in row])
    worksheet.auto_filter.ref = (
        f"A1:{get_column_letter(len(columns))}{len(frame) + 1}"
    )


def write_excel(path: Path, stocks: pd.DataFrame, indexes: pd.DataFrame) -> None:
    stock_view, index_view = excel_view_frames(stocks, indexes)
    workbook = Workbook(write_only=True)
    _append_sheet(workbook, "stocks", stock_view, STOCK_COLUMNS)
    _append_sheet(workbook, "indexes", index_view, INDEX_COLUMNS)
    workbook.save(path)


def verify_excel(path: Path, stocks: pd.DataFrame, indexes: pd.DataFrame) -> None:
    stock_view, index_view = excel_view_frames(stocks, indexes)
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if workbook.sheetnames != ["stocks", "indexes"]:
            raise RuntimeError(
                f"Excel sheets are {workbook.sheetnames}, expected ['stocks', 'indexes']"
            )
        expected = {
            "stocks": (stock_view, STOCK_COLUMNS),
            "indexes": (index_view, INDEX_COLUMNS),
        }
        for name, (frame, columns) in expected.items():
            worksheet = workbook[name]
            rows = worksheet.iter_rows(values_only=True)
            header = tuple(next(rows))
            if header != columns:
                raise RuntimeError(f"Excel {name} header does not match the schema")
            date_column = columns.index("trade_date") + 1
            dates: set[str] = set()
            row_count = 0
            for row in rows:
                row_count += 1
                dates.add(str(row[date_column - 1]))
            if row_count != len(frame):
                raise RuntimeError(
                    f"Excel {name} has {row_count} data rows, expected {len(frame)}"
                )
            if row_count > EXCEL_MAX_DATA_ROWS:
                raise RuntimeError(f"Excel {name} exceeds the worksheet row limit")
            expected_dates = set(frame["trade_date"].astype(str))
            if dates != expected_dates:
                raise RuntimeError(f"Excel {name} contains incomplete date boundaries")
            if max(dates) != str(frame["trade_date"].max()):
                raise RuntimeError(f"Excel {name} does not include the latest date")
    finally:
        workbook.close()


def write_bundle_outputs(data_root: Path, bundle: dict[str, Any]) -> tuple[Path, Path]:
    validated = validate_bundle(bundle)
    pickle_path, excel_path = bundle_paths(data_root)
    pickle_path.parent.mkdir(parents=True, exist_ok=True)
    nonce = uuid4().hex
    pickle_temporary = pickle_path.with_name(f"market_data.{nonce}.tmp.pkl")
    excel_temporary = excel_path.with_name(f"market_data.{nonce}.tmp.xlsx")
    try:
        pd.to_pickle(validated, pickle_temporary)
        reloaded = load_bundle(pickle_temporary)
        if not reloaded["stocks"].equals(validated["stocks"]):
            raise RuntimeError("Temporary Pickle changed stocks content or data types")
        if not reloaded["indexes"].equals(validated["indexes"]):
            raise RuntimeError("Temporary Pickle changed indexes content or data types")

        write_excel(excel_temporary, reloaded["stocks"], reloaded["indexes"])
        verify_excel(excel_temporary, reloaded["stocks"], reloaded["indexes"])

        os.replace(pickle_temporary, pickle_path)
        os.replace(excel_temporary, excel_path)
    finally:
        for temporary in (pickle_temporary, excel_temporary):
            if temporary.exists():
                temporary.unlink()
    return pickle_path, excel_path


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Add missing Tushare trading days to market_data.pkl and Excel."
    )
    parser.add_argument(
        "--start",
        type=parse_trade_date,
        help=(
            "Inclusive start date. Required only when market_data.pkl does not yet "
            "exist; otherwise the earliest bundled date is used."
        ),
    )
    parser.add_argument(
        "--end",
        type=parse_trade_date,
        help=(
            "Inclusive end date. Defaults to today after 18:00 China time, otherwise "
            "yesterday; weekends and holidays are removed by the exchange calendar."
        ),
    )
    parser.add_argument(
        "--data-root",
        type=Path,
        default=DEFAULT_DATA_ROOT,
        help="Data directory containing market_data.pkl and market_data.xlsx.",
    )
    parser.add_argument(
        "--timeout",
        type=int,
        default=30,
        help="Tushare HTTP timeout in seconds. Defaults to 30.",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    data_root = args.data_root.resolve()
    pickle_path, _ = bundle_paths(data_root)
    if pickle_path.is_file():
        bundle = load_bundle(pickle_path)
        stocks: pd.DataFrame | None = bundle["stocks"]
        indexes: pd.DataFrame | None = bundle["indexes"]
        existing_dates = sorted(set(stocks["trade_date"].astype(str)))
        start_date = args.start or existing_dates[0]
    else:
        if args.start is None:
            raise RuntimeError(
                "market_data.pkl does not exist. Supply --start for the initial "
                "Tushare backfill."
            )
        stocks = None
        indexes = None
        existing_dates = []
        start_date = args.start
    end_date = args.end or default_end_date()

    load_dotenv(PROJECT_ROOT / ".env", override=False)
    token = os.getenv("TUSHARE_TOKEN")
    if not token:
        raise RuntimeError("TUSHARE_TOKEN is missing from the project .env file")

    new_stocks: list[pd.DataFrame] = []
    new_indexes: list[pd.DataFrame] = []
    if start_date <= end_date:
        print(
            f"[Tushare] SDK {ts.__version__}; checking {start_date} through {end_date}",
            flush=True,
        )
        api = ts.pro_api(token, timeout=args.timeout)
        open_dates = query_open_dates(api, start_date, end_date)
        missing_dates = missing_trade_dates(open_dates, existing_dates)
        for position, trade_date in enumerate(missing_dates, start=1):
            print(
                f"[Tushare] {position}/{len(missing_dates)} downloading {trade_date}",
                flush=True,
            )
            frames = fetch_trade_date(api, trade_date)
            new_stocks.append(build_stocks_frame(frames, trade_date))
            new_indexes.append(build_indexes_frame(frames, trade_date))
    elif args.start is not None:
        raise ValueError(f"Start date {start_date} is after end date {end_date}")

    if new_stocks:
        stock_parts = new_stocks if stocks is None else [stocks, *new_stocks]
        index_parts = new_indexes if indexes is None else [indexes, *new_indexes]
        stocks = pd.concat(stock_parts, ignore_index=True)
        indexes = pd.concat(index_parts, ignore_index=True)
        stocks = stocks.sort_values(["trade_date", "ts_code"], kind="stable").reset_index(
            drop=True
        )
        indexes = indexes.sort_values(
            ["trade_date", "ts_code"], kind="stable"
        ).reset_index(drop=True)
        print(f"[Tushare] adding {len(new_stocks)} complete trading days", flush=True)
    else:
        if stocks is None or indexes is None:
            raise RuntimeError(
                f"No open trading days were found from {start_date} through {end_date}; "
                "the initial bundle was not created."
            )
        print("[Tushare] no missing trading days; rebuilding Excel from Pickle", flush=True)

    output_bundle = make_bundle(stocks, indexes)
    final_pickle, final_excel = write_bundle_outputs(data_root, output_bundle)
    print(f"[Tushare] validated Pickle written to {final_pickle}", flush=True)
    print(f"[Tushare] validated Excel written to {final_excel}", flush=True)


if __name__ == "__main__":
    main()
